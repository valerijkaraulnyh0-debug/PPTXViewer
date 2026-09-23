# -*- coding: utf-8 -*-
"""
PPTX Viewer for Android — просмотр PowerPoint презентаций.
Парсит .pptx через python-pptx, рендерит слайды в Kivy.
Навигация: свайп влево/вправо или кнопки.
"""
import os
from kivy.app import App
from kivy.uix.boxlayout import BoxLayout
from kivy.uix.floatlayout import FloatLayout
from kivy.uix.label import Label
from kivy.uix.image import Image
from kivy.uix.button import Button
from kivy.uix.popup import Popup
from kivy.graphics import Color, Rectangle
from kivy.properties import ObjectProperty, StringProperty, NumericProperty

from pptx import Presentation
from pptx.util import Emu


class SlideWidget(FloatLayout):
    """Виджет одного слайда."""
    def __init__(self, slide, slide_w, slide_h, **kw):
        super().__init__(**kw)
        self.slide = slide
        self.slide_w = slide_w
        self.slide_h = slide_h
        self.bind(size=self._layout, pos=self._layout)

    def _layout(self, *a):
        self.canvas.clear()
        w, h = self.size
        if w <= 0 or h <= 0:
            return
        scale = min(w / self.slide_w, h / self.slide_h)
        ox = (w - self.slide_w * scale) / 2
        oy = (h - self.slide_h * scale) / 2

        with self.canvas:
            Color(1, 1, 1, 1)
            Rectangle(pos=(ox, oy), size=(self.slide_w * scale, self.slide_h * scale))

        for child in list(self.children):
            self.remove_widget(child)

        for shape in self.slide.shapes:
            sx = ox + (shape.left or 0) * scale
            sy = oy + (self.slide_h - (shape.top or 0) - (shape.height or 0)) * scale
            sw = (shape.width or 100) * scale
            sh = (shape.height or 50) * scale

            if shape.has_text_frame:
                txt = shape.text_frame.text.strip()
                if txt:
                    font_size = max(10, min(28, int(14 * scale)))
                    lbl = Label(
                        text=txt, font_size=font_size, color=(0, 0, 0, 1),
                        size_hint=(None, None), size=(sw, sh), pos=(sx, sy),
                        halign="left", valign="top", markup=True, text_size=(sw, None),
                    )
                    self.add_widget(lbl)

            if shape.shape_type == 13:
                try:
                    img_bytes = shape.image.blob
                    tmp = os.path.join("/tmp", f"slide_img_{id(shape)}.png")
                    with open(tmp, "wb") as f:
                        f.write(img_bytes)
                    img = Image(
                        source=tmp, size_hint=(None, None), size=(sw, sh),
                        pos=(sx, sy), fit_mode="contain",
                    )
                    self.add_widget(img)
                except Exception:
                    pass


class PPTXViewerApp(App):
    current_slide = NumericProperty(0)
    total_slides = NumericProperty(0)
    file_path = StringProperty("")

    def build(self):
        self.title = "PPTX Viewer"
        layout = BoxLayout(orientation="vertical")

        top_bar = BoxLayout(size_hint=(1, 0.08))
        self.btn_prev = Button(text="< Назад", size_hint=(0.3, 1))
        self.slide_label = Label(text="Слайд 0 / 0", size_hint=(0.4, 1))
        self.btn_next = Button(text="Вперёд >", size_hint=(0.3, 1))
        self.btn_prev.bind(on_press=self.prev_slide)
        self.btn_next.bind(on_press=self.next_slide)
        top_bar.add_widget(self.btn_prev)
        top_bar.add_widget(self.slide_label)
        top_bar.add_widget(self.btn_next)
        layout.add_widget(top_bar)

        self.slide_area = FloatLayout(size_hint=(1, 0.84))
        layout.add_widget(self.slide_area)

        bottom_bar = BoxLayout(size_hint=(1, 0.08))
        self.btn_open = Button(text="Открыть файл", size_hint=(0.5, 1))
        self.btn_open.bind(on_press=self.open_file_dialog)
        bottom_bar.add_widget(self.btn_open)
        layout.add_widget(bottom_bar)

        self._touch_start = None
        layout.bind(on_touch_down=self.on_touch_down)
        layout.bind(on_touch_up=self.on_touch_up)
        self.presentation = None
        return layout

    def on_touch_down(self, widget, touch):
        self._touch_start = (touch.x, touch.y)
        return False

    def on_touch_up(self, widget, touch):
        if self._touch_start is None:
            return False
        dx = touch.x - self._touch_start[0]
        dy = touch.y - self._touch_start[1]
        self._touch_start = None
        if abs(dx) > 50 and abs(dx) > abs(dy):
            if dx > 0:
                self.prev_slide()
            else:
                self.next_slide()
        return False

    def open_file_dialog(self, *a):
        try:
            from android.permissions import request_permissions, Permission
            request_permissions([Permission.READ_EXTERNAL_STORAGE])
            from jnius import autoclass
            Intent = autoclass("android.content.Intent")
            PythonActivity = autoclass("org.kivy.android.PythonActivity")
            intent = Intent(Intent.ACTION_GET_CONTENT)
            intent.setType("application/vnd.openxmlformats-officedocument.presentationml.presentation")
            intent.addCategory(Intent.CATEGORY_OPENABLE)
            PythonActivity.mActivity.startActivityForResult(
                Intent.createChooser(intent, "Choose PPTX"), 42)
        except ImportError:
            # Десктопный режим — файловый диалог Kivy
            from kivy.uix.filechooser import FileChooserListView
            content = BoxLayout(orientation="vertical")
            fc = FileChooserListView(filters=["*.pptx"])
            content.add_widget(fc)
            btn = Button(text="Open", size_hint=(1, 0.1))
            content.add_widget(btn)
            popup = Popup(title="Open PPTX", content=content, size_hint=(0.9, 0.9))
            btn.bind(on_press=lambda x: self._desktop_open(fc.selection, popup))
            popup.open()

    def _desktop_open(self, selection, popup):
        popup.dismiss()
        if selection:
            self.load_pptx(selection[0])

    def on_pause(self):
        return True

    def on_resume(self):
        pass

    def load_pptx(self, path):
        self.file_path = path
        try:
            self.presentation = Presentation(path)
        except Exception as e:
            popup = Popup(title="Error",
                          content=Label(text=f"Cannot open:\n{e}"),
                          size_hint=(0.8, 0.4))
            popup.open()
            return
        self.total_slides = len(self.presentation.slides)
        self.current_slide = 0
        self.render_slide()

    def render_slide(self):
        if not self.presentation or self.total_slides == 0:
            return
        self.slide_area.clear_widgets()
        slide = self.presentation.slides[self.current_slide]
        sw = self.presentation.slide_width or 9144000
        sh = self.presentation.slide_height or 5143500
        sw_pt = Emu(sw).pt
        sh_pt = Emu(sh).pt
        sw_px = int(sw_pt * 96 / 72)
        sh_px = int(sh_pt * 96 / 72)
        sw_widget = SlideWidget(slide, sw_px, sh_px, size_hint=(1, 1))
        self.slide_area.add_widget(sw_widget)
        self.slide_label.text = f"Слайд {self.current_slide + 1} / {self.total_slides}"
        self.btn_prev.disabled = self.current_slide == 0
        self.btn_next.disabled = self.current_slide == self.total_slides - 1

    def prev_slide(self, *a):
        if self.current_slide > 0:
            self.current_slide -= 1
            self.render_slide()

    def next_slide(self, *a):
        if self.current_slide < self.total_slides - 1:
            self.current_slide += 1
            self.render_slide()


if __name__ == "__main__":
    PPTXViewerApp().run()
